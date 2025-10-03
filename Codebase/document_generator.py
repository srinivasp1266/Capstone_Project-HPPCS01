"""
Document generation system for creating formatted CV documents.
"""

import os
import io
from pathlib import Path
from typing import Dict, Any, Optional, List
from datetime import datetime
import logging

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_LINE_SPACING
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml.ns import nsdecls
    from docx.oxml import parse_xml
except ImportError:
    Document = None

try:
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_JUSTIFY
except ImportError:
    canvas = None
    SimpleDocTemplate = None

try:
    from jinja2 import Template
except ImportError:
    Template = None

try:
    import pdfkit
except ImportError:
    pdfkit = None

try:
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
    from reportlab.lib.units import inch
    from reportlab.lib.colors import black, blue, grey
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False

logger = logging.getLogger(__name__)


class CVDocumentGenerator:
    """Generator for creating formatted CV documents."""
    
    def __init__(self, templates_dir: str = "templates"):
        """
        Initialize document generator.
        
        Args:
            templates_dir: Directory containing template files
        """
        self.templates_dir = Path(templates_dir)
        self.templates_dir.mkdir(parents=True, exist_ok=True)
        self._create_default_templates()
    
    def _create_default_templates(self):
        """Create default template files if they don't exist."""
        # Create a simple HTML template for web preview
        html_template_path = self.templates_dir / "cv_template.html"
        if not html_template_path.exists():
            self._create_html_template(html_template_path)
        
        # Create a markdown template
        md_template_path = self.templates_dir / "cv_template.md"
        if not md_template_path.exists():
            self._create_markdown_template(md_template_path)
    
    def _create_html_template(self, template_path: Path):
        """Create default HTML template."""
        html_template = """
<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{{ personal_info.name or "Professional Resume" }}</title>
    <style>
        body {
            font-family: 'Arial', sans-serif;
            line-height: 1.6;
            margin: 0;
            padding: 20px;
            color: #333;
            max-width: 800px;
            margin: 0 auto;
        }
        
        .header {
            text-align: center;
            border-bottom: 2px solid #2c3e50;
            padding-bottom: 20px;
            margin-bottom: 30px;
        }
        
        .name {
            font-size: 2.5em;
            font-weight: bold;
            color: #2c3e50;
            margin-bottom: 10px;
        }
        
        .contact-info {
            font-size: 1.1em;
            color: #7f8c8d;
            margin-bottom: 5px;
        }
        
        .section {
            margin-bottom: 30px;
        }
        
        .section-title {
            font-size: 1.4em;
            font-weight: bold;
            color: #2c3e50;
            border-bottom: 1px solid #bdc3c7;
            padding-bottom: 5px;
            margin-bottom: 15px;
            text-transform: uppercase;
        }
        
        .experience-item, .education-item, .project-item {
            margin-bottom: 20px;
            padding-left: 20px;
            border-left: 3px solid #3498db;
        }
        
        .position-title, .degree-title, .project-title {
            font-weight: bold;
            font-size: 1.2em;
            color: #2c3e50;
        }
        
        .company, .institution {
            font-weight: bold;
            color: #3498db;
            font-size: 1.1em;
        }
        
        .dates {
            color: #7f8c8d;
            font-style: italic;
        }
        
        .location {
            color: #95a5a6;
            float: right;
        }
        
        .summary {
            font-size: 1.1em;
            line-height: 1.8;
            text-align: justify;
            font-style: italic;
            color: #34495e;
            background: #f8f9fa;
            padding: 20px;
            border-left: 4px solid #3498db;
            margin-bottom: 30px;
        }
        
        .skills-grid {
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(250px, 1fr));
            gap: 15px;
        }
        
        .skill-category {
            background: #f8f9fa;
            padding: 15px;
            border-radius: 5px;
            border-left: 4px solid #3498db;
        }
        
        .skill-category h4 {
            margin-top: 0;
            color: #2c3e50;
            text-transform: uppercase;
            font-size: 1em;
        }
        
        .skill-list {
            margin: 0;
            padding: 0;
            list-style: none;
        }
        
        .skill-list li {
            display: inline-block;
            background: #3498db;
            color: white;
            padding: 4px 12px;
            margin: 2px;
            border-radius: 15px;
            font-size: 0.9em;
        }
        
        .responsibility, .achievement {
            margin-left: 20px;
            margin-bottom: 5px;
        }
        
        .responsibility::before, .achievement::before {
            content: "▸ ";
            color: #3498db;
            font-weight: bold;
        }
        
        .technologies {
            margin-top: 10px;
        }
        
        .tech-tag {
            display: inline-block;
            background: #95a5a6;
            color: white;
            padding: 2px 8px;
            margin: 2px;
            border-radius: 10px;
            font-size: 0.8em;
        }
        
        @media print {
            body { font-size: 12pt; }
            .section { page-break-inside: avoid; }
        }
    </style>
</head>
<body>
    <div class="header">
        <div class="name">{{ personal_info.name or "Your Name" }}</div>
        {% if personal_info.email %}
        <div class="contact-info">📧 {{ personal_info.email }}</div>
        {% endif %}
        {% if personal_info.phone %}
        <div class="contact-info">📞 {{ personal_info.phone }}</div>
        {% endif %}
        {% if personal_info.address %}
        <div class="contact-info">📍 {{ personal_info.address }}</div>
        {% endif %}
        {% if personal_info.linkedin %}
        <div class="contact-info">🔗 <a href="{{ personal_info.linkedin }}">LinkedIn</a></div>
        {% endif %}
        {% if personal_info.github %}
        <div class="contact-info">🔗 <a href="{{ personal_info.github }}">GitHub</a></div>
        {% endif %}
    </div>

    {% if summary %}
    <div class="section">
        <div class="summary">{{ summary }}</div>
    </div>
    {% endif %}

    {% if experience %}
    <div class="section">
        <div class="section-title">Professional Experience</div>
        {% for exp in experience %}
        <div class="experience-item">
            <div class="position-title">{{ exp.position or "Position Title" }}</div>
            <div class="company">{{ exp.company or "Company Name" }}
                {% if exp.location %}<span class="location">{{ exp.location }}</span>{% endif %}
            </div>
            <div class="dates">{{ exp.start_date or "Start Date" }} - {{ exp.end_date or "End Date" }}</div>
            
            {% if exp.responsibilities %}
            <div style="margin-top: 10px;">
                {% for resp in exp.responsibilities %}
                <div class="responsibility">{{ resp }}</div>
                {% endfor %}
            </div>
            {% endif %}
            
            {% if exp.achievements %}
            <div style="margin-top: 10px;">
                {% for achievement in exp.achievements %}
                <div class="achievement">{{ achievement }}</div>
                {% endfor %}
            </div>
            {% endif %}
            
            {% if exp.technologies %}
            <div class="technologies">
                {% for tech in exp.technologies %}
                <span class="tech-tag">{{ tech }}</span>
                {% endfor %}
            </div>
            {% endif %}
        </div>
        {% endfor %}
    </div>
    {% endif %}

    {% if education %}
    <div class="section">
        <div class="section-title">Education</div>
        {% for edu in education %}
        <div class="education-item">
            <div class="degree-title">{{ edu.degree or "Degree" }}</div>
            <div class="institution">{{ edu.institution or "Institution" }}
                {% if edu.location %}<span class="location">{{ edu.location }}</span>{% endif %}
            </div>
            <div class="dates">{{ edu.start_date or "Start" }} - {{ edu.end_date or "End" }}</div>
            {% if edu.gpa %}
            <div>GPA: {{ edu.gpa }}</div>
            {% endif %}
            
            {% if edu.relevant_coursework %}
            <div style="margin-top: 10px;">
                <strong>Relevant Coursework:</strong>
                {% for course in edu.relevant_coursework %}
                <span class="tech-tag">{{ course }}</span>
                {% endfor %}
            </div>
            {% endif %}
        </div>
        {% endfor %}
    </div>
    {% endif %}

    {% if skills and (skills.technical or skills.soft or skills.tools) %}
    <div class="section">
        <div class="section-title">Skills</div>
        <div class="skills-grid">
            {% if skills.technical %}
            <div class="skill-category">
                <h4>Technical Skills</h4>
                <ul class="skill-list">
                    {% for skill in skills.technical %}
                    <li>{{ skill }}</li>
                    {% endfor %}
                </ul>
            </div>
            {% endif %}
            
            {% if skills.tools %}
            <div class="skill-category">
                <h4>Tools & Technologies</h4>
                <ul class="skill-list">
                    {% for skill in skills.tools %}
                    <li>{{ skill }}</li>
                    {% endfor %}
                </ul>
            </div>
            {% endif %}
            
            {% if skills.soft %}
            <div class="skill-category">
                <h4>Soft Skills</h4>
                <ul class="skill-list">
                    {% for skill in skills.soft %}
                    <li>{{ skill }}</li>
                    {% endfor %}
                </ul>
            </div>
            {% endif %}
        </div>
    </div>
    {% endif %}

    {% if projects %}
    <div class="section">
        <div class="section-title">Projects</div>
        {% for project in projects %}
        <div class="project-item">
            <div class="project-title">{{ project.name or "Project Name" }}</div>
            <div class="dates">{{ project.start_date or "" }}{% if project.start_date and project.end_date %} - {% endif %}{{ project.end_date or "" }}</div>
            
            {% if project.description %}
            <div style="margin-top: 10px;">{{ project.description }}</div>
            {% endif %}
            
            {% if project.technologies %}
            <div class="technologies" style="margin-top: 10px;">
                {% for tech in project.technologies %}
                <span class="tech-tag">{{ tech }}</span>
                {% endfor %}
            </div>
            {% endif %}
            
            {% if project.url or project.github_url %}
            <div style="margin-top: 10px;">
                {% if project.url %}<a href="{{ project.url }}">Live Demo</a>{% endif %}
                {% if project.url and project.github_url %} | {% endif %}
                {% if project.github_url %}<a href="{{ project.github_url }}">GitHub</a>{% endif %}
            </div>
            {% endif %}
        </div>
        {% endfor %}
    </div>
    {% endif %}

    {% if certifications %}
    <div class="section">
        <div class="section-title">Certifications</div>
        {% for cert in certifications %}
        <div class="experience-item">
            <div class="position-title">{{ cert.name or "Certification Name" }}</div>
            <div class="company">{{ cert.issuer or "Issuer" }}</div>
            <div class="dates">{{ cert.date or "Date" }}</div>
        </div>
        {% endfor %}
    </div>
    {% endif %}

    {% if achievements %}
    <div class="section">
        <div class="section-title">Achievements</div>
        {% for achievement in achievements %}
        <div class="responsibility">{{ achievement }}</div>
        {% endfor %}
    </div>
    {% endif %}
</body>
</html>
        """
        
        with open(template_path, 'w', encoding='utf-8') as f:
            f.write(html_template.strip())
    
    def _create_markdown_template(self, template_path: Path):
        """Create default Markdown template."""
        md_template = """
# {{ personal_info.name or "Your Name" }}

{% if personal_info.email or personal_info.phone or personal_info.address %}
**Contact Information:**
{% if personal_info.email %}📧 {{ personal_info.email }}{% endif %}
{% if personal_info.phone %}📞 {{ personal_info.phone }}{% endif %}
{% if personal_info.address %}📍 {{ personal_info.address }}{% endif %}
{% if personal_info.linkedin %}🔗 [LinkedIn]({{ personal_info.linkedin }}){% endif %}
{% if personal_info.github %}🔗 [GitHub]({{ personal_info.github }}){% endif %}

---
{% endif %}

{% if summary %}
## Professional Summary

{{ summary }}

---
{% endif %}

{% if experience %}
## Professional Experience

{% for exp in experience %}
### {{ exp.position or "Position Title" }}
**{{ exp.company or "Company Name" }}** {% if exp.location %}| {{ exp.location }}{% endif %}  
*{{ exp.start_date or "Start Date" }} - {{ exp.end_date or "End Date" }}*

{% if exp.responsibilities %}
{% for resp in exp.responsibilities %}
- {{ resp }}
{% endfor %}
{% endif %}

{% if exp.achievements %}
**Key Achievements:**
{% for achievement in exp.achievements %}
- {{ achievement }}
{% endfor %}
{% endif %}

{% if exp.technologies %}
**Technologies:** {{ exp.technologies | join(', ') }}
{% endif %}

{% endfor %}
---
{% endif %}

{% if education %}
## Education

{% for edu in education %}
### {{ edu.degree or "Degree" }}
**{{ edu.institution or "Institution" }}** {% if edu.location %}| {{ edu.location }}{% endif %}  
*{{ edu.start_date or "Start" }} - {{ edu.end_date or "End" }}*
{% if edu.gpa %}  
**GPA:** {{ edu.gpa }}{% endif %}

{% if edu.relevant_coursework %}
**Relevant Coursework:** {{ edu.relevant_coursework | join(', ') }}
{% endif %}

{% endfor %}
---
{% endif %}

{% if skills and (skills.technical or skills.soft or skills.tools) %}
## Skills

{% if skills.technical %}
**Technical Skills:** {{ skills.technical | join(', ') }}
{% endif %}

{% if skills.tools %}
**Tools & Technologies:** {{ skills.tools | join(', ') }}
{% endif %}

{% if skills.soft %}
**Soft Skills:** {{ skills.soft | join(', ') }}
{% endif %}

---
{% endif %}

{% if projects %}
## Projects

{% for project in projects %}
### {{ project.name or "Project Name" }}
{% if project.start_date or project.end_date %}*{{ project.start_date or "" }}{% if project.start_date and project.end_date %} - {% endif %}{{ project.end_date or "" }}*{% endif %}

{% if project.description %}
{{ project.description }}
{% endif %}

{% if project.technologies %}
**Technologies:** {{ project.technologies | join(', ') }}
{% endif %}

{% if project.url or project.github_url %}
{% if project.url %}[Live Demo]({{ project.url }}){% endif %}{% if project.url and project.github_url %} | {% endif %}{% if project.github_url %}[GitHub]({{ project.github_url }}){% endif %}
{% endif %}

{% endfor %}
---
{% endif %}

{% if certifications %}
## Certifications

{% for cert in certifications %}
- **{{ cert.name or "Certification Name" }}** - {{ cert.issuer or "Issuer" }} ({{ cert.date or "Date" }})
{% endfor %}

---
{% endif %}

{% if achievements %}
## Achievements

{% for achievement in achievements %}
- {{ achievement }}
{% endfor %}
{% endif %}
        """
        
        with open(template_path, 'w', encoding='utf-8') as f:
            f.write(md_template.strip())
    
    def generate_html(self, cv_data: Dict[str, Any], template_name: str = "cv_template.html") -> str:
        """
        Generate HTML document from CV data.
        
        Args:
            cv_data: CV data dictionary
            template_name: Name of the template file
            
        Returns:
            Generated HTML content
        """
        if Template is None:
            raise ImportError("Jinja2 is required for HTML generation. Install with: pip install jinja2")
        
        template_path = self.templates_dir / template_name
        
        if not template_path.exists():
            self._create_html_template(template_path)
        
        with open(template_path, 'r', encoding='utf-8') as f:
            template_content = f.read()
        
        template = Template(template_content)
        return template.render(**cv_data)
    
    def generate_pdf_from_html(self, cv_data: Dict[str, Any], template_name: str = "cv_template.html") -> bytes:
        """
        Generate PDF from CV data using ReportLab (cross-platform solution).
        
        Args:
            cv_data: CV data dictionary
            template_name: Name of the template file (for compatibility)
            
        Returns:
            Generated PDF content as bytes
        """
        if not REPORTLAB_AVAILABLE:
            raise ImportError("ReportLab is required for PDF generation. Install with: pip install reportlab")
        
        from io import BytesIO
        buffer = BytesIO()
        
        # Create PDF document
        doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=1*inch)
        story = []
        
        # Get styles
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Title'],
            fontSize=24,
            spaceAfter=12,
            textColor=blue
        )
        heading_style = ParagraphStyle(
            'CustomHeading',
            parent=styles['Heading2'],
            fontSize=14,
            spaceAfter=6,
            textColor=black,
            fontName='Helvetica-Bold'
        )
        normal_style = styles['Normal']
        
        # Extract data - matching HTML template structure
        personal_info = cv_data.get('personal_info', {})
        experience = cv_data.get('experience', [])  # Changed from work_experience
        education = cv_data.get('education', [])
        skills = cv_data.get('skills', [])
        projects = cv_data.get('projects', [])
        certifications = cv_data.get('certifications', [])
        achievements = cv_data.get('achievements', [])
        summary = cv_data.get('summary', '')  # Top-level summary
        
        # Add personal information
        name = personal_info.get('name', 'Professional Resume')
        story.append(Paragraph(name, title_style))
        
        contact_details = []
        if personal_info.get('email'):
            contact_details.append(f"📧 {personal_info['email']}")
        if personal_info.get('phone'):
            contact_details.append(f"📞 {personal_info['phone']}")
        if personal_info.get('address'):
            contact_details.append(f"📍 {personal_info['address']}")
        if personal_info.get('linkedin'):
            contact_details.append(f"🔗 LinkedIn: {personal_info['linkedin']}")
        if personal_info.get('github'):
            contact_details.append(f"🔗 GitHub: {personal_info['github']}")
            
        if contact_details:
            story.append(Paragraph(" | ".join(contact_details), normal_style))
        story.append(Spacer(1, 0.2*inch))
        
        # Add summary if available (top-level summary)
        if summary:
            story.append(Paragraph("Professional Summary", heading_style))
            story.append(Paragraph(summary, normal_style))
            story.append(Spacer(1, 0.2*inch))
        
        # Add work experience - matching HTML template structure
        if experience:
            story.append(Paragraph("Professional Experience", heading_style))
            for exp in experience:
                if isinstance(exp, dict):
                    # Position and company
                    exp_title = f"<b>{exp.get('position', 'Position Title')}</b>"
                    company_info = exp.get('company', 'Company Name')
                    if exp.get('location'):
                        company_info += f" ({exp['location']})"
                    story.append(Paragraph(exp_title, normal_style))
                    story.append(Paragraph(company_info, normal_style))
                    
                    # Dates
                    start_date = exp.get('start_date', 'Start Date')
                    end_date = exp.get('end_date', 'End Date')
                    story.append(Paragraph(f"<i>{start_date} - {end_date}</i>", normal_style))
                    
                    # Responsibilities
                    if exp.get('responsibilities'):
                        for resp in exp['responsibilities']:
                            story.append(Paragraph(f"▸ {resp}", normal_style))
                    
                    story.append(Spacer(1, 0.15*inch))
        
        # Add education
        if education:
            story.append(Paragraph("Education", heading_style))
            for edu in education:
                if isinstance(edu, dict):
                    edu_title = f"<b>{edu.get('degree', 'Degree')}</b>"
                    story.append(Paragraph(edu_title, normal_style))
                    if edu.get('institution'):
                        story.append(Paragraph(edu['institution'], normal_style))
                    if edu.get('year') or edu.get('graduation_date'):
                        date_text = edu.get('year') or edu.get('graduation_date')
                        story.append(Paragraph(f"<i>{date_text}</i>", normal_style))
                    story.append(Spacer(1, 0.1*inch))
        
        # Add skills
        if skills:
            story.append(Paragraph("Skills", heading_style))
            if isinstance(skills, dict):
                # Handle categorized skills
                for category, skill_list in skills.items():
                    story.append(Paragraph(f"<b>{category}:</b>", normal_style))
                    if isinstance(skill_list, list):
                        skills_text = ", ".join(skill_list)
                    else:
                        skills_text = str(skill_list)
                    story.append(Paragraph(skills_text, normal_style))
                    story.append(Spacer(1, 0.05*inch))
            else:
                # Handle flat skill list
                skills_text = ", ".join([skill if isinstance(skill, str) else str(skill) for skill in skills])
                story.append(Paragraph(skills_text, normal_style))
            story.append(Spacer(1, 0.15*inch))
        
        # Add projects
        if projects:
            story.append(Paragraph("Projects", heading_style))
            for project in projects:
                if isinstance(project, dict):
                    project_title = f"<b>{project.get('name', 'Project Name')}</b>"
                    story.append(Paragraph(project_title, normal_style))
                    
                    if project.get('description'):
                        story.append(Paragraph(project['description'], normal_style))
                    
                    if project.get('technologies'):
                        tech_text = "Technologies: " + ", ".join(project['technologies'])
                        story.append(Paragraph(f"<i>{tech_text}</i>", normal_style))
                    
                    links = []
                    if project.get('url'):
                        links.append(f"Live Demo: {project['url']}")
                    if project.get('github_url'):
                        links.append(f"GitHub: {project['github_url']}")
                    if links:
                        story.append(Paragraph(" | ".join(links), normal_style))
                    
                    story.append(Spacer(1, 0.1*inch))
        
        # Add certifications
        if certifications:
            story.append(Paragraph("Certifications", heading_style))
            for cert in certifications:
                if isinstance(cert, dict):
                    cert_text = f"<b>{cert.get('name', 'Certification Name')}</b>"
                    if cert.get('issuer'):
                        cert_text += f" - {cert['issuer']}"
                    story.append(Paragraph(cert_text, normal_style))
                    if cert.get('date'):
                        story.append(Paragraph(f"<i>{cert['date']}</i>", normal_style))
                    story.append(Spacer(1, 0.1*inch))
        
        # Add achievements
        if achievements:
            story.append(Paragraph("Achievements", heading_style))
            for achievement in achievements:
                story.append(Paragraph(f"▸ {achievement}", normal_style))
        
        # Build PDF
        doc.build(story)
        pdf_bytes = buffer.getvalue()
        buffer.close()
        
        return pdf_bytes
    
    def generate_markdown(self, cv_data: Dict[str, Any], template_name: str = "cv_template.md") -> str:
        """
        Generate Markdown document from CV data.
        
        Args:
            cv_data: CV data dictionary
            template_name: Name of the template file
            
        Returns:
            Generated Markdown content
        """
        if Template is None:
            raise ImportError("Jinja2 is required for Markdown generation. Install with: pip install jinja2")
        
        template_path = self.templates_dir / template_name
        
        if not template_path.exists():
            self._create_markdown_template(template_path)
        
        with open(template_path, 'r', encoding='utf-8') as f:
            template_content = f.read()
        
        template = Template(template_content)
        return template.render(**cv_data)
    
    def generate_docx(self, cv_data: Dict[str, Any], output_path: Optional[str] = None) -> str:
        """
        Generate Word document from CV data.
        
        Args:
            cv_data: CV data dictionary
            output_path: Output file path (optional)
            
        Returns:
            Path to generated document
        """
        if Document is None:
            raise ImportError("python-docx is required for DOCX generation. Install with: pip install python-docx")
        
        doc = Document()
        
        # Set document margins
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)
        
        # Define styles
        self._create_docx_styles(doc)
        
        # Add header with personal information
        self._add_docx_header(doc, cv_data.get('personal_info', {}))
        
        # Add professional summary
        if cv_data.get('summary'):
            self._add_docx_summary(doc, cv_data['summary'])
        
        # Add experience section
        if cv_data.get('experience'):
            self._add_docx_experience(doc, cv_data['experience'])
        
        # Add education section
        if cv_data.get('education'):
            self._add_docx_education(doc, cv_data['education'])
        
        # Add skills section
        if cv_data.get('skills'):
            self._add_docx_skills(doc, cv_data['skills'])
        
        # Add projects section
        if cv_data.get('projects'):
            self._add_docx_projects(doc, cv_data['projects'])
        
        # Add certifications section
        if cv_data.get('certifications'):
            self._add_docx_certifications(doc, cv_data['certifications'])
        
        # Add achievements section
        if cv_data.get('achievements'):
            self._add_docx_achievements(doc, cv_data['achievements'])
        
        # Save document
        if not output_path:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            name = cv_data.get('personal_info', {}).get('name', 'CV').replace(' ', '_')
            output_path = f"data/output/{name}_{timestamp}.docx"
        
        # Ensure output directory exists
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        
        doc.save(output_path)
        return output_path
    
    def generate_pdf(self, cv_data: Dict[str, Any], output_path: Optional[str] = None) -> str:
        """
        Generate PDF document from CV data.
        
        Args:
            cv_data: CV data dictionary
            output_path: Output file path (optional)
            
        Returns:
            Path to generated document
        """
        if SimpleDocTemplate is None:
            raise ImportError("reportlab is required for PDF generation. Install with: pip install reportlab")
        
        if not output_path:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            name = cv_data.get('personal_info', {}).get('name', 'CV').replace(' ', '_')
            output_path = f"data/output/{name}_{timestamp}.pdf"
        
        # Ensure output directory exists
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        
        # Create PDF document
        doc = SimpleDocTemplate(output_path, pagesize=letter)
        story = []
        
        # Define styles
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle('Title', parent=styles['Title'], fontSize=24, spaceAfter=12)
        heading_style = ParagraphStyle('Heading', parent=styles['Heading2'], fontSize=14, spaceAfter=6)
        normal_style = styles['Normal']
        
        # Add header
        personal_info = cv_data.get('personal_info', {})
        if personal_info.get('name'):
            story.append(Paragraph(personal_info['name'], title_style))
        
        # Contact information
        contact_info = []
        if personal_info.get('email'):
            contact_info.append(f"Email: {personal_info['email']}")
        if personal_info.get('phone'):
            contact_info.append(f"Phone: {personal_info['phone']}")
        if personal_info.get('address'):
            contact_info.append(f"Address: {personal_info['address']}")
        
        if contact_info:
            story.append(Paragraph(" | ".join(contact_info), normal_style))
            story.append(Spacer(1, 12))
        
        # Professional summary
        if cv_data.get('summary'):
            story.append(Paragraph("PROFESSIONAL SUMMARY", heading_style))
            story.append(Paragraph(cv_data['summary'], normal_style))
            story.append(Spacer(1, 12))
        
        # Experience
        if cv_data.get('experience'):
            story.append(Paragraph("PROFESSIONAL EXPERIENCE", heading_style))
            for exp in cv_data['experience']:
                # Position and company
                exp_title = f"<b>{exp.get('position', 'Position')}</b> - {exp.get('company', 'Company')}"
                if exp.get('location'):
                    exp_title += f" ({exp['location']})"
                story.append(Paragraph(exp_title, normal_style))
                
                # Dates
                dates = f"{exp.get('start_date', 'Start')} - {exp.get('end_date', 'End')}"
                story.append(Paragraph(f"<i>{dates}</i>", normal_style))
                
                # Responsibilities
                if exp.get('responsibilities'):
                    for resp in exp['responsibilities']:
                        story.append(Paragraph(f"• {resp}", normal_style))
                
                story.append(Spacer(1, 6))
        
        # Build PDF
        doc.build(story)
        return output_path
    
    def _create_docx_styles(self, doc):
        """Create custom styles for the Word document."""
        styles = doc.styles
        
        # Header style
        if 'Header' not in [style.name for style in styles]:
            header_style = styles.add_style('Header', WD_STYLE_TYPE.PARAGRAPH)
            header_style.font.size = Pt(18)
            header_style.font.bold = True
            header_style.font.color.rgb = RGBColor(44, 62, 80)
            header_style.paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            header_style.paragraph_format.space_after = Pt(12)
        
        # Section heading style
        if 'SectionHeading' not in [style.name for style in styles]:
            section_style = styles.add_style('SectionHeading', WD_STYLE_TYPE.PARAGRAPH)
            section_style.font.size = Pt(14)
            section_style.font.bold = True
            section_style.font.color.rgb = RGBColor(44, 62, 80)
            section_style.paragraph_format.space_before = Pt(12)
            section_style.paragraph_format.space_after = Pt(6)
            
            # Add bottom border
            p = section_style.element
            pPr = p.get_or_add_pPr()
            pBdr = parse_xml(r'<w:pBdr %s><w:bottom w:val="single" w:sz="4" w:space="1" w:color="BDC3C7"/></w:pBdr>' % nsdecls('w'))
            pPr.insert_element_before(pBdr, 'w:shd', 'w:tabs', 'w:suppressAutoHyphens', 'w:kinsoku', 'w:wordWrap')
    
    def _add_docx_header(self, doc, personal_info):
        """Add header section to Word document."""
        # Name
        name_para = doc.add_paragraph(personal_info.get('name', 'Your Name'), style='Header')
        
        # Contact information
        contact_parts = []
        if personal_info.get('email'):
            contact_parts.append(personal_info['email'])
        if personal_info.get('phone'):
            contact_parts.append(personal_info['phone'])
        if personal_info.get('address'):
            contact_parts.append(personal_info['address'])
        
        if contact_parts:
            contact_para = doc.add_paragraph(' | '.join(contact_parts))
            contact_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        
        # LinkedIn and GitHub
        if personal_info.get('linkedin') or personal_info.get('github'):
            links = []
            if personal_info.get('linkedin'):
                links.append(personal_info['linkedin'])
            if personal_info.get('github'):
                links.append(personal_info['github'])
            
            links_para = doc.add_paragraph(' | '.join(links))
            links_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    
    def _add_docx_summary(self, doc, summary):
        """Add professional summary to Word document."""
        doc.add_paragraph('PROFESSIONAL SUMMARY', style='SectionHeading')
        doc.add_paragraph(summary)
    
    def _add_docx_experience(self, doc, experiences):
        """Add experience section to Word document."""
        doc.add_paragraph('PROFESSIONAL EXPERIENCE', style='SectionHeading')
        
        for exp in experiences:
            # Position and company
            exp_para = doc.add_paragraph()
            exp_para.add_run(exp.get('position', 'Position')).bold = True
            exp_para.add_run(f" - {exp.get('company', 'Company')}")
            
            if exp.get('location'):
                location_run = exp_para.add_run(f" ({exp['location']})")
                location_run.font.color.rgb = RGBColor(149, 165, 166)
            
            # Dates
            dates = f"{exp.get('start_date', 'Start')} - {exp.get('end_date', 'End')}"
            date_para = doc.add_paragraph(dates)
            date_para.runs[0].italic = True
            
            # Responsibilities
            if exp.get('responsibilities'):
                for resp in exp['responsibilities']:
                    bullet_para = doc.add_paragraph(f"• {resp}")
                    bullet_para.paragraph_format.left_indent = Inches(0.25)
    
    def _add_docx_education(self, doc, educations):
        """Add education section to Word document."""
        doc.add_paragraph('EDUCATION', style='SectionHeading')
        
        for edu in educations:
            edu_para = doc.add_paragraph()
            edu_para.add_run(edu.get('degree', 'Degree')).bold = True
            edu_para.add_run(f" - {edu.get('institution', 'Institution')}")
            
            dates = f"{edu.get('start_date', 'Start')} - {edu.get('end_date', 'End')}"
            date_para = doc.add_paragraph(dates)
            date_para.runs[0].italic = True
            
            if edu.get('gpa'):
                doc.add_paragraph(f"GPA: {edu['gpa']}")
    
    def _add_docx_skills(self, doc, skills):
        """Add skills section to Word document."""
        doc.add_paragraph('SKILLS', style='SectionHeading')
        
        if skills.get('technical'):
            tech_para = doc.add_paragraph()
            tech_para.add_run('Technical Skills: ').bold = True
            tech_para.add_run(', '.join(skills['technical']))
        
        if skills.get('tools'):
            tools_para = doc.add_paragraph()
            tools_para.add_run('Tools & Technologies: ').bold = True
            tools_para.add_run(', '.join(skills['tools']))
        
        if skills.get('soft'):
            soft_para = doc.add_paragraph()
            soft_para.add_run('Soft Skills: ').bold = True
            soft_para.add_run(', '.join(skills['soft']))
    
    def _add_docx_projects(self, doc, projects):
        """Add projects section to Word document."""
        doc.add_paragraph('PROJECTS', style='SectionHeading')
        
        for project in projects:
            proj_para = doc.add_paragraph()
            proj_para.add_run(project.get('name', 'Project Name')).bold = True
            
            if project.get('description'):
                doc.add_paragraph(project['description'])
            
            if project.get('technologies'):
                tech_para = doc.add_paragraph()
                tech_para.add_run('Technologies: ').bold = True
                tech_para.add_run(', '.join(project['technologies']))
    
    def _add_docx_certifications(self, doc, certifications):
        """Add certifications section to Word document."""
        doc.add_paragraph('CERTIFICATIONS', style='SectionHeading')
        
        for cert in certifications:
            cert_text = cert.get('name', 'Certification')
            if cert.get('issuer'):
                cert_text += f" - {cert['issuer']}"
            if cert.get('date'):
                cert_text += f" ({cert['date']})"
            
            doc.add_paragraph(f"• {cert_text}")
    
    def _add_docx_achievements(self, doc, achievements):
        """Add achievements section to Word document."""
        doc.add_paragraph('ACHIEVEMENTS', style='SectionHeading')
        
        for achievement in achievements:
            doc.add_paragraph(f"• {achievement}")
    
    def save_to_file(self, content: str, output_path: str, file_format: str = "html") -> str:
        """
        Save content to file.
        
        Args:
            content: Content to save
            output_path: Output file path
            file_format: File format (html, md, txt)
            
        Returns:
            Path to saved file
        """
        # Ensure output directory exists
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(content)
        
        return output_path
    
    def get_available_templates(self) -> List[str]:
        """Get list of available template files."""
        templates = []
        for file_path in self.templates_dir.glob("*.html"):
            templates.append(file_path.name)
        for file_path in self.templates_dir.glob("*.md"):
            templates.append(file_path.name)
        
        return templates